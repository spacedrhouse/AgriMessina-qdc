import pandas as pd
from sqlalchemy import text
from PyQt6.QtWidgets import (QVBoxLayout, QHBoxLayout, QPushButton, QMessageBox,
                             QDialog, QComboBox, QLabel, QLineEdit, QDoubleSpinBox,
                             QSpinBox, QFormLayout, QScrollArea, QWidget, QDateEdit,
                             QTableView, QFileDialog, QHeaderView)
from PyQt6.QtCore import Qt, QDate
from PyQt6.QtGui import QStandardItemModel, QStandardItem

from ui_core import PannelloBaseDialog, DelegateMovimenti
from app_logging import get_logger

log = get_logger(__name__)


# Conversione fra UM compatibili. Storiamo le qta nel registro nel numeratore
# di `unita_misura` (l'UM "interna" del prodotto), ma l'utente inserisce i
# carichi nella sua `unita_carico`. Al salvataggio convertiamo, alla lettura
# riconvertiamo.
_MASSA_AL_KG = {"mg": 1e-6, "g": 1e-3, "kg": 1.0}
_VOLUME_AL_L = {"ml": 1e-3, "l": 1.0}


def converti_qta(qta: float, da_um: str | None, a_um: str | None) -> float:
    """Converte `qta` da `da_um` a `a_um`. Identità se UM uguali, vuote o
    sconosciute (es. "Unità"), o se le due UM appartengono a categorie diverse
    (massa ↔ volume non è una conversione valida).

    Esempi: converti_qta(1, "kg", "g") → 1000; converti_qta(500, "ml", "l") → 0.5.
    """
    if not da_um or not a_um or qta is None:
        return qta
    da = str(da_um).strip().lower()
    a = str(a_um).strip().lower()
    if da == a:
        return qta
    if da in _MASSA_AL_KG and a in _MASSA_AL_KG:
        return qta * _MASSA_AL_KG[da] / _MASSA_AL_KG[a]
    if da in _VOLUME_AL_L and a in _VOLUME_AL_L:
        return qta * _VOLUME_AL_L[da] / _VOLUME_AL_L[a]
    return qta

class DialogProdotto(QDialog):
    CAMPI = [
        ("nome_prodotto",            "Nome Prodotto *",            "text",         None),
        ("categoria",                "Categoria",                  "text",         None),
        ("numero_registrazione",     "N° Registrazione",           "text",         None),
        ("sostanza_attiva",          "Sostanza Attiva",            "text",         None),
        ("bio_convenzionale",        "Bio / Convenzionale",        "combo",        ["", "Bio", "Conv"]),
        ("avversita",                "Avversità",                  "text",         None),
        ("titolo_n",                 "Titolo N",                   "decimal",      None),
        ("titolo_p",                 "Titolo P",                   "decimal",      None),
        ("titolo_k",                 "Titolo K",                   "decimal",      None),
        ("phi_giorni",               "PHI (giorni)",               "intero",       None),
        ("trattamenti_max",          "Max Trattamenti/Anno",       "intero",       None),
        ("intervallo_min_tratt",     "Intervallo Minimo (giorni)", "intero",       None),
        ("unita_misura",             "Unità di Misura",            "combo",        ["", "L/ha", "kg/ha", "ml/ha", "g/ha", "g/hl", "ml/hl", "unità/ha"]),
        ("unita_carico",             "Unità di Carico",            "combo",        [""]),
        ("min_sostanza",             "Dose Minima",                "decimal",      None),
        ("max_sostanza",             "Dose Massima",               "decimal",      None),
        ("qta_acqua",                "Quantità Acqua (L/ha)",      "decimal",      None),
        ("blacklist",                "Blacklist",                  "combo",        ["No", "Si"]),
    ]

    @staticmethod
    def opzioni_unita_carico(unita_misura: str | None) -> list[str]:
        """UM ammesse per i carichi manuali in funzione del numeratore di
        `unita_misura`. Lista vuota se l'UM non è ancora stata scelta."""
        if not unita_misura:
            return []
        num = str(unita_misura).split('/')[0].strip().lower()
        if num in ('mg', 'g', 'kg'):
            return ['mg', 'g', 'kg']
        if num in ('ml', 'l'):
            return ['ml', 'l']
        if num == 'unità':
            return ['Unità']
        return []

    def __init__(self, engine, dati=None, parent=None):
        super().__init__(parent)
        self.engine = engine
        self.dati   = dati or {}
        self.widget_map = {}

        self.setWindowTitle("Modifica Prodotto" if dati else "Nuovo Prodotto")
        self.setMinimumWidth(480)

        outer  = QVBoxLayout(self)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        contenuto = QWidget()
        form = QFormLayout(contenuto)
        form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        form.setSpacing(8)
        scroll.setWidget(contenuto)
        outer.addWidget(scroll)

        for campo, etichetta, tipo, opzioni in self.CAMPI:
            valore = self.dati.get(campo)
            if tipo == "text":
                w = QLineEdit()
                w.setText(str(valore) if valore is not None else "")
            elif tipo in ("combo", "combo_cascata", "combo_modale"):
                w = QComboBox()
                w.addItems(opzioni)
                if valore and str(valore) in opzioni:
                    w.setCurrentText(str(valore))
            elif tipo == "decimal":
                w = QDoubleSpinBox()
                w.setRange(0.0, 999999.99)
                w.setDecimals(2)
                w.setSpecialValueText("—")
                try: w.setValue(float(valore)) if valore is not None else w.setValue(0.0)
                except: w.setValue(0.0)
            elif tipo == "intero":
                w = QSpinBox()
                w.setRange(0, 999999)
                w.setSpecialValueText("—")
                try: w.setValue(int(valore)) if valore is not None else w.setValue(0)
                except: w.setValue(0)
            self.widget_map[campo] = (w, tipo)
            form.addRow(etichetta, w)

        # Cablaggio dinamico: la lista di "Unità di Carico" dipende dal
        # numeratore di "Unità di Misura". Cambiandola, ripopoliamo l'altro
        # combo (preservando il valore corrente se ancora valido).
        um_w, _ = self.widget_map["unita_misura"]
        uc_w, _ = self.widget_map["unita_carico"]
        um_w.currentTextChanged.connect(self._aggiorna_opzioni_unita_carico)
        # Popolamento iniziale: usa il valore di unita_misura corrente
        # (in apertura il combo è già stato settato sopra).
        self._aggiorna_opzioni_unita_carico(um_w.currentText(),
                                            preserva=self.dati.get("unita_carico"))

        btns = QHBoxLayout()
        btn_salva   = QPushButton("💾 Salva")
        btn_annulla = QPushButton("Annulla")

        btn_salva.setProperty('class', 'success')
        btn_annulla.setProperty('class', 'secondary')

        btn_salva.clicked.connect(self.salva)
        btn_annulla.clicked.connect(self.reject)
        btns.addWidget(btn_salva)
        btns.addWidget(btn_annulla)
        outer.addLayout(btns)

    def _aggiorna_opzioni_unita_carico(self, unita_misura: str,
                                        preserva: str | None = None) -> None:
        """Ripopola il combo "Unità di Carico" in base al numeratore dell'UM.
        Preserva la selezione corrente se ancora valida nella nuova lista."""
        uc_w, _ = self.widget_map["unita_carico"]
        valore_corrente = preserva if preserva is not None else uc_w.currentText().strip()

        opzioni = [""] + self.opzioni_unita_carico(unita_misura)

        uc_w.blockSignals(True)
        uc_w.clear()
        uc_w.addItems(opzioni)
        if valore_corrente and valore_corrente in opzioni:
            uc_w.setCurrentText(valore_corrente)
        else:
            # Niente match: default sull'unità "standard" (l'ultima della lista,
            # tipicamente kg per massa, l per volume, Unità per unità).
            if len(opzioni) > 1:
                uc_w.setCurrentText(opzioni[-1])
        uc_w.blockSignals(False)

    @staticmethod
    def _numeratore_um(um: str | None) -> str | None:
        if not um:
            return None
        return str(um).split('/')[0].strip().lower()

    def salva(self):
        valori = {}
        for campo, (w, tipo) in self.widget_map.items():
            if tipo == "text": v = w.text().strip() or None
            elif tipo in ("combo", "combo_cascata", "combo_modale"): v = w.currentText().strip() or None
            elif tipo == "decimal": v = w.value() if w.value() > 0.0 else None
            elif tipo == "intero": v = w.value() if w.value() > 0 else None
            valori[campo] = v

        if not valori.get("nome_prodotto"):
            QMessageBox.critical(self, "Errore", "Il campo 'Nome Prodotto' è obbligatorio!")
            return

        # Guardia: cambio del numeratore di `unita_misura` su un prodotto che
        # ha già carichi/scarichi manuali. Le qta nel registro sono storate
        # nel numeratore vecchio; con il nuovo numeratore vengono interpretate
        # in modo diverso (es. 10 "kg" → 10 "g" = 1000x meno sostanza), falsando
        # dose cumulativa e giacenze. Avvisa esplicitamente prima di salvare.
        if self.dati.get("id"):
            old_num = self._numeratore_um(self.dati.get("unita_misura"))
            new_num = self._numeratore_um(valori.get("unita_misura"))
            if old_num and new_num and old_num != new_num:
                with self.engine.connect() as conn:
                    n_mov = conn.execute(text(
                        "SELECT COUNT(*) FROM registro_magazzino "
                        "WHERE prodotto_id = :pid AND trattamento_id IS NULL"
                    ), {"pid": self.dati["id"]}).scalar() or 0
                if n_mov > 0:
                    risposta = QMessageBox.warning(
                        self, "Attenzione: cambio unità di misura",
                        f"Stai cambiando l'unità di misura del prodotto da "
                        f"<b>{self.dati.get('unita_misura')}</b> a "
                        f"<b>{valori.get('unita_misura')}</b>.<br><br>"
                        f"Esistono <b>{n_mov}</b> carichi/scarichi manuali storici "
                        f"per questo prodotto, registrati nell'UM <b>{old_num}</b>. "
                        f"Il loro valore numerico resta invariato ma sarà ora "
                        f"interpretato come <b>{new_num}</b>, falsando dosi e "
                        f"giacenze (es. 10 {old_num} verranno letti come 10 {new_num}).<br><br>"
                        f"Prima di procedere considera di convertire manualmente i "
                        f"valori storici, oppure di rimuoverli.<br><br>"
                        f"Procedere comunque?",
                        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                        QMessageBox.StandardButton.No,
                    )
                    if risposta != QMessageBox.StandardButton.Yes:
                        return

        try:
            # I trigger su prodotti accodano automaticamente l'operazione in
            # pending_operations: niente enqueue manuale qui, altrimenti
            # finirebbe in due pending op e quindi in due record server-side.
            with self.engine.begin() as conn:
                if self.dati.get("id"):
                    set_clause = ", ".join(f"{c} = :{c}" for c in valori)
                    valori["id"] = self.dati["id"]
                    conn.execute(text(f"UPDATE prodotti SET {set_clause} WHERE id = :id"), valori)
                else:
                    cols = ", ".join(valori.keys())
                    ph   = ", ".join(f":{c}" for c in valori)
                    # Se questo execute fallisce (es. nome duplicato),
                    # salta all'except e NON salva nulla.
                    conn.execute(text(f"INSERT INTO prodotti ({cols}) VALUES ({ph})"), valori)
            self.accept()

        except Exception as e:
            # Se siamo qui, il blocco 'with' ha fatto il ROLLBACK automatico in locale
            messaggio = str(e)
            if "UNIQUE constraint failed" in messaggio:
                QMessageBox.warning(self, "Prodotto Esistente",
                                    f"Errore: Il prodotto '{valori['nome_prodotto']}' esiste già.\n"
                                    "L'operazione è stata annullata e non è stato salvato nulla.")
            else:
                QMessageBox.critical(self, "Errore Database", f"Salvataggio fallito: {messaggio}")

class DialogNuovoMovimento(QDialog):
    def __init__(self, engine, prodotto_id, um, parent=None, unita_carico=None):
        super().__init__(parent)
        self.engine = engine
        self.prodotto_id = prodotto_id
        # `um`: numeratore di unita_misura del prodotto (es. "g"); è l'unità
        # interna in cui il registro_magazzino storerà la quantità.
        # `unita_carico`: unità con cui l'utente inserisce (es. "kg"). Se None
        #  o uguale a `um`, nessuna conversione viene applicata.
        self.um_interna = str(um or "")
        self.unita_carico = str(unita_carico or "").strip() or self.um_interna
        self.setWindowTitle("Registra Movimento Magazzino")
        self.setMinimumWidth(400)

        layout = QVBoxLayout(self)
        form = QFormLayout()
        self.data_edit = QDateEdit(QDate.currentDate())
        self.data_edit.setCalendarPopup(True)
        self.combo_tipo = QComboBox()
        self.combo_tipo.addItems(["CARICO", "SCARICO"])

        self.combo_azienda = QComboBox()
        self.combo_azienda.addItem("Nessuna / Magazzino Generale", userData=None)
        with self.engine.connect() as conn:
            for az in conn.execute(text("SELECT id, nome FROM aziende ORDER BY nome")).fetchall():
                self.combo_azienda.addItem(az[1], userData=az[0])

        self.spin_qta = QDoubleSpinBox()
        # Min 1e-6 (= 1 mg in kg, 1 µl in l): carichi piccoli convertiti
        # in UM grandi (es. 5 g salvati come 0.005 kg) altrimenti verrebbero
        # clippati al minimo del widget e ri-salvati come valore errato.
        self.spin_qta.setRange(0.000001, 999999.99)
        self.spin_qta.setDecimals(4)
        self.spin_qta.setSuffix(f" {self.unita_carico}")

        # Etichetta informativa: mostra il valore convertito nell'UM interna,
        # così l'utente capisce in quale unità il registro lo storerà.
        self.lbl_convertito = QLabel("")
        self.lbl_convertito.setStyleSheet("color: #757575; font-size: 11px;")
        self.spin_qta.valueChanged.connect(self._aggiorna_anteprima_conversione)

        self.edit_ddt = QLineEdit()
        self.edit_fornitore = QLineEdit()
        self.edit_note = QLineEdit()

        form.addRow("Data Movimento:", self.data_edit)
        form.addRow("Tipo Operazione:", self.combo_tipo)
        form.addRow("Azienda Proprietaria:", self.combo_azienda)
        form.addRow("Quantità:", self.spin_qta)
        if self.unita_carico.lower() != self.um_interna.lower() and self.um_interna:
            form.addRow("", self.lbl_convertito)
            self._aggiorna_anteprima_conversione(self.spin_qta.value())
        form.addRow("N° DDT (Opzionale):", self.edit_ddt)
        form.addRow("Fornitore (Opzionale):", self.edit_fornitore)
        form.addRow("Note:", self.edit_note)
        layout.addLayout(form)

        btns = QHBoxLayout()
        btn_salva = QPushButton("💾 Registra Movimento")
        btn_salva.setProperty('class', 'success')
        btn_salva.clicked.connect(self.salva)
        btns.addWidget(btn_salva)
        layout.addLayout(btns)

    def _aggiorna_anteprima_conversione(self, valore: float) -> None:
        convertito = converti_qta(valore, self.unita_carico, self.um_interna)
        self.lbl_convertito.setText(
            f"≡ {convertito:.4f} {self.um_interna} nel registro"
        )

    def salva(self):
        try:
            # Arrotondo a 6 decimali: senza, ogni open-save degrada per
            # imprecisione float (es. 0.7 kg -> 699.999...g, salvato, riaperto
            # -> 0.6999... kg, salvato, ecc). 6 decimali sono sufficienti per
            # tutte le UM supportate (mg = 1e-6 kg).
            qta_storage = round(converti_qta(self.spin_qta.value(),
                                             self.unita_carico, self.um_interna), 6)
            with self.engine.begin() as conn:
                conn.execute(text("""
                    INSERT INTO registro_magazzino (prodotto_id, azienda_id, data_movimento, tipo_movimento, quantita, n_ddt, fornitore, note)
                    VALUES (:pid, :az, :d, :tipo, :q, :ddt, :fornitore, :note)
                """), {
                    "pid": self.prodotto_id, "az": self.combo_azienda.currentData(),
                    "d": self.data_edit.date().toPyDate(), "tipo": self.combo_tipo.currentText(),
                    "q": qta_storage, "ddt": self.edit_ddt.text().strip() or None,
                    "fornitore": self.edit_fornitore.text().strip() or None, "note": self.edit_note.text().strip() or None
                })
            self.accept()
        except Exception as e: QMessageBox.critical(self, "Errore", str(e))


class DialogModificaMovimento(QDialog):
    def __init__(self, engine, dati, um, parent=None, unita_carico=None):
        super().__init__(parent)
        self.engine = engine
        self.id_mov = dati['id']
        # Stessa semantica di DialogNuovoMovimento: l'UI parla in `unita_carico`,
        # il DB stora in `um_interna`. La qta che arriva in `dati['qta']` è già
        # in `um_interna` (letta dal registro), quindi va convertita per la
        # visualizzazione e ri-convertita al salvataggio.
        self.um_interna = str(um or "")
        self.unita_carico = str(unita_carico or "").strip() or self.um_interna
        self.setWindowTitle("Modifica Movimento Magazzino")

        layout = QVBoxLayout(self)
        form = QFormLayout()

        self.data_edit = QDateEdit()
        self.data_edit.setCalendarPopup(True)
        self.data_edit.setDate(QDate.fromString(dati['data'], Qt.DateFormat.ISODate) if isinstance(dati['data'], str) else dati['data'])

        self.combo_tipo = QComboBox()
        self.combo_tipo.addItems(["CARICO", "SCARICO"])
        self.combo_tipo.setCurrentText(dati['tipo'])
        self.combo_tipo.setEnabled(False)

        self.combo_azienda = QComboBox()
        self.combo_azienda.addItem("Nessuna / Magazzino Generale", userData=None)
        with self.engine.connect() as conn:
            for az in conn.execute(text("SELECT id, nome FROM aziende ORDER BY nome")).fetchall():
                self.combo_azienda.addItem(az[1], userData=az[0])
                if az[0] == dati['azienda_id']: self.combo_azienda.setCurrentIndex(self.combo_azienda.count() - 1)

        self.spin_qta = QDoubleSpinBox()
        # Min 1e-6 (= 1 mg in kg, 1 µl in l): carichi piccoli convertiti
        # in UM grandi (es. 5 g salvati come 0.005 kg) altrimenti verrebbero
        # clippati al minimo del widget e ri-salvati come valore errato.
        self.spin_qta.setRange(0.000001, 999999.99)
        self.spin_qta.setDecimals(4)
        self.spin_qta.setSuffix(f" {self.unita_carico}")
        # Converte la qta storata (in um_interna) verso unita_carico per
        # mostrarla nell'UM con cui l'utente l'aveva inserita.
        self.spin_qta.setValue(converti_qta(float(dati['qta']),
                                            self.um_interna, self.unita_carico))

        self.lbl_convertito = QLabel("")
        self.lbl_convertito.setStyleSheet("color: #757575; font-size: 11px;")
        self.spin_qta.valueChanged.connect(self._aggiorna_anteprima_conversione)

        self.edit_ddt = QLineEdit(str(dati['ddt'] or ""))
        self.edit_fornitore = QLineEdit(str(dati['fornitore'] or ""))
        self.edit_note = QLineEdit(str(dati['note'] or ""))

        form.addRow("Data Movimento:", self.data_edit)
        form.addRow("Tipo Operazione:", self.combo_tipo)
        form.addRow("Azienda Proprietaria:", self.combo_azienda)
        form.addRow("Quantità:", self.spin_qta)
        if self.unita_carico.lower() != self.um_interna.lower() and self.um_interna:
            form.addRow("", self.lbl_convertito)
            self._aggiorna_anteprima_conversione(self.spin_qta.value())
        form.addRow("N° DDT:", self.edit_ddt)
        form.addRow("Fornitore:", self.edit_fornitore)
        form.addRow("Note:", self.edit_note)
        layout.addLayout(form)

        btn_salva = QPushButton("💾 Salva Modifiche")
        btn_salva.setProperty('class', 'warning')
        btn_salva.clicked.connect(self.salva)
        layout.addWidget(btn_salva)

    def _aggiorna_anteprima_conversione(self, valore: float) -> None:
        convertito = converti_qta(valore, self.unita_carico, self.um_interna)
        self.lbl_convertito.setText(
            f"≡ {convertito:.4f} {self.um_interna} nel registro"
        )

    def salva(self):
        try:
            # Arrotondo a 6 decimali: senza, ogni open-save degrada per
            # imprecisione float (es. 0.7 kg -> 699.999...g, salvato, riaperto
            # -> 0.6999... kg, salvato, ecc). 6 decimali sono sufficienti per
            # tutte le UM supportate (mg = 1e-6 kg).
            qta_storage = round(converti_qta(self.spin_qta.value(),
                                             self.unita_carico, self.um_interna), 6)
            with self.engine.begin() as conn:
                conn.execute(text("""
                    UPDATE registro_magazzino SET azienda_id=:az, data_movimento=:d, tipo_movimento=:tipo,
                    quantita=:q, n_ddt=:ddt, fornitore=:fornitore, note=:note WHERE id=:id
                """), {
                    "az": self.combo_azienda.currentData(), "d": self.data_edit.date().toPyDate(),
                    "tipo": self.combo_tipo.currentText(), "q": qta_storage,
                    "ddt": self.edit_ddt.text().strip() or None, "fornitore": self.edit_fornitore.text().strip() or None,
                    "note": self.edit_note.text().strip() or None, "id": self.id_mov
                })
            self.accept()
        except Exception as e: QMessageBox.critical(self, "Errore", str(e))


class DialogRegistroProdotto(QDialog):
    def __init__(self, engine, db, prodotto_id, nome_prodotto, um, parent=None,
                 azienda_filter: str | None = None, azienda_ids: list | None = None,
                 mostra_fittizio: bool = False):
        super().__init__(parent)
        self.engine, self.db = engine, db
        self.prodotto_id, self.um = prodotto_id, str(um or "")
        self.azienda_filter = azienda_filter
        self.azienda_ids = azienda_ids or []
        # In modalità fittizio il dialog è READ-ONLY: il fittizio è alimentato
        # solo automaticamente (dai trattamenti revisionati). Nessuna scrittura
        # manuale è permessa.
        self.mostra_fittizio = mostra_fittizio
        self.tabella = "registro_magazzino_fittizio" if mostra_fittizio \
            else "registro_magazzino"

        # --- ESTRAIAMO L'UNITÀ DI MISURA ASSOLUTA (es. da "kg/ha" a "kg") ---
        self.um_pulita = self.um.split('/')[0].strip() if '/' in self.um else self.um

        # `unita_carico`: UM con cui l'utente inserisce i carichi manuali.
        # I dialog di nuovo/modifica la useranno come unità di input e
        # convertiranno in `um_pulita` (UM "interna" del registro) al salvataggio.
        with self.engine.connect() as conn:
            self.unita_carico = conn.execute(text(
                "SELECT unita_carico FROM prodotti WHERE id = :pid"
            ), {"pid": self.prodotto_id}).scalar()

        suffisso = " — Fittizio" if mostra_fittizio else ""
        titolo_filter = f" ({azienda_filter})" if azienda_filter else ""
        self.setWindowTitle(
            f"Registro Carico/Scarico — {nome_prodotto}{titolo_filter}{suffisso}"
        )
        self.setMinimumSize(950, 500)

        layout = QVBoxLayout(self)
        self.lbl_giacenza = QLabel()
        self.lbl_giacenza.setStyleSheet("font-size: 18px; font-weight: bold; padding: 10px; border-bottom: 1px solid #555;")
        layout.addWidget(self.lbl_giacenza)

        h_tool = QHBoxLayout()
        btn_nuovo = QPushButton("➕ Nuovo Movimento")
        btn_modifica = QPushButton("✏️ Modifica")
        btn_elimina = QPushButton("🗑️ Elimina")

        btn_nuovo.setProperty('class', 'success')
        btn_modifica.setProperty('class', 'warning')
        btn_elimina.setProperty('class', 'danger')

        btn_nuovo.clicked.connect(self._nuovo_movimento)
        btn_modifica.clicked.connect(self._modifica_movimento)
        btn_elimina.clicked.connect(self._elimina_movimento)
        # Read-only in vista fittizio: i bottoni di scrittura manuale sono
        # nascosti del tutto (il fittizio si popola solo dai trattamenti
        # revisionati, quindi non c'è nulla da fare a mano).
        if self.mostra_fittizio:
            for b in (btn_nuovo, btn_modifica, btn_elimina):
                b.setVisible(False)
        h_tool.addWidget(btn_nuovo); h_tool.addWidget(btn_modifica); h_tool.addWidget(btn_elimina); h_tool.addStretch()
        layout.addLayout(h_tool)

        self.vista = QTableView()
        self.vista.setSelectionBehavior(QTableView.SelectionBehavior.SelectRows)
        self.vista.setItemDelegate(DelegateMovimenti(self))

        header = self.vista.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        header.setStretchLastSection(True)

        layout.addWidget(self.vista)
        self.aggiorna_dati()

    def aggiorna_dati(self):
        # Clausola per filtrare per magazzino (se passato dal pannello padre)
        clausola_az = ""
        if self.azienda_ids:
            ids_str = ",".join(str(int(i)) for i in self.azienda_ids)
            clausola_az = f"AND rm.azienda_id IN ({ids_str})"

        with self.engine.connect() as conn:
            giacenza = conn.execute(text(f"""
                SELECT SUM(CASE WHEN rm.tipo_movimento = 'CARICO' THEN rm.quantita ELSE -rm.quantita END)
                FROM {self.tabella} rm
                WHERE rm.prodotto_id = :pid {clausola_az}
            """), {"pid": self.prodotto_id}).scalar() or 0.0

        label_giacenza = f"Giacenza ({self.azienda_filter})" if self.azienda_filter else "Giacenza Attuale"
        self.lbl_giacenza.setText(
            f"{label_giacenza}: <span style='color:{'green' if giacenza>=0 else 'red'}'>"
            f"{giacenza:.4f} {self.um_pulita}</span>"
        )

        # La colonna "Azienda" mostra l'azienda DEL TENDONE originale (pre-alias):
        # COALESCE(az_orig, az_warehouse, 'Generale').
        # Esempio: per Deflorio Ciccopinto via Messina Alfio, mostra "Deflorio Ciccopinto"
        # (più informativo che "Messina Alfio" che è il solo magazzino).
        # Per i CARICHI manuali (azienda_id_origine NULL) fallback su azienda_id.
        query_sql = f"""
            SELECT rm.id, rm.data_movimento AS "Data", rm.tipo_movimento AS "Tipo",
                   COALESCE(az_orig.nome, az_wh.nome, 'Generale') AS "Azienda",
                   rm.quantita AS "Quantità ({self.um_pulita})",
                   rm.n_ddt AS "N. DDT", rm.fornitore AS "Fornitore",
                   rm.note AS "Note", rm.azienda_id
            FROM {self.tabella} rm
            LEFT JOIN aziende az_wh ON az_wh.id = rm.azienda_id
            LEFT JOIN aziende az_orig ON az_orig.id = rm.azienda_id_origine
            WHERE rm.prodotto_id = :pid {clausola_az}
            ORDER BY rm.data_movimento DESC, rm.id DESC
        """

        # ---> INIZIO SOSTITUZIONE CON QStandardItemModel <---
        self.modello_registro = QStandardItemModel()

        with self.engine.connect() as conn:
            result = conn.execute(text(query_sql), {"pid": self.prodotto_id})
            col_names = list(result.keys())
            self.modello_registro.setHorizontalHeaderLabels(col_names)

            for row in result:
                items = []
                row_dict = dict(row._mapping)
                for val in row:
                    item = QStandardItem(str(val) if val is not None else "")
                    item.setEditable(False)
                    items.append(item)
                items[0].setData(row_dict, Qt.ItemDataRole.UserRole)
                self.modello_registro.appendRow(items)

        self.vista.setModel(self.modello_registro)
        self.vista.setColumnHidden(0, True)
        self.vista.setColumnHidden(8, True)
        self.vista.resizeColumnsToContents()

    # Per far funzionare il "finto record" anche nei due metodi sottostanti
    def _recupera_finto_record(self, idx):
        item = self.modello_registro.item(idx.row(), 0)
        dati = item.data(Qt.ItemDataRole.UserRole)
        class DummyRecord:
            def __init__(self, d): self.d = d
            def value(self, k): return self.d.get(k)
        return DummyRecord(dati)

    def _nuovo_movimento(self):
        # `unita_carico` se definita guida l'UM di input nel dialog (con
        # conversione automatica verso `um_pulita` al salvataggio).
        if DialogNuovoMovimento(self.engine, self.prodotto_id, self.um_pulita,
                                self, unita_carico=self.unita_carico).exec():
            self.aggiorna_dati()

    def _modifica_movimento(self):
        idx = self.vista.currentIndex()
        if not idx.isValid(): return
        r = self._recupera_finto_record(idx) # <--- MODIFICATO QUI

        if r.value("Tipo") == "SCARICO":
            QMessageBox.warning(self, "Bloccato", "Gli scarichi non possono essere modificati qui.")
            return

        dati = {"id": r.value("id"), "data": r.value("Data"), "tipo": r.value("Tipo"), "qta": float(r.value(f"Quantità ({self.um_pulita})")), "ddt": r.value("N. DDT"), "fornitore": r.value("Fornitore"), "note": r.value("Note"), "azienda_id": r.value("azienda_id")}
        if DialogModificaMovimento(self.engine, dati, self.um_pulita,
                                   self, unita_carico=self.unita_carico).exec():
            self.aggiorna_dati()

    def _elimina_movimento(self):
        idx = self.vista.currentIndex()
        if not idx.isValid(): return
        r = self._recupera_finto_record(idx) # <--- MODIFICATO QUI

        if QMessageBox.question(self, "Conferma", "Vuoi eliminare?") == QMessageBox.StandardButton.Yes:
            with self.engine.begin() as conn: conn.execute(text("DELETE FROM registro_magazzino WHERE id=:id"), {"id": r.value("id")})
            self.aggiorna_dati()


class PannelloProdotti(PannelloBaseDialog):
    COLONNE_NASCOSTE = [0]

    def __init__(self, engine, db, azienda_filter: str):
        """
        Args:
            azienda_filter: nome dell'azienda magazzino (es. "Agrimessina").
                Il pannello mostra solo i movimenti/giacenze di quell'azienda
                + degli alias che mappano su di essa (vedi WAREHOUSE_ALIASES).
        """
        super().__init__()
        self.engine, self.db = engine, db
        self.azienda_filter = azienda_filter
        self.azienda_ids = self._resolve_filter_ids()  # lista per IN clause SQL

        # Toggle reale ↔ fittizio. False = reale (default), True = fittizio.
        # In modalità fittizio i bottoni di modifica manuale sono disabilitati:
        # il fittizio è popolato solo automaticamente dai trattamenti revisionati.
        self.mostra_fittizio = False

        self.btn_export_tutti_mov = QPushButton("📊 Esporta Movimenti")
        self.btn_export_tutti_mov.setProperty('class', 'success')
        self.btn_export_tutti_mov.clicked.connect(self._esporta_tutti_movimenti)

        # Esporta giacenze: snapshot prodotto-per-prodotto con SUM(CARICO) -
        # SUM(SCARICO). Rispetta il filtro azienda e il toggle reale/fittizio
        # selezionato. Utile per inventario fisico / verifiche etichetta.
        self.btn_export_giacenze = QPushButton("📦 Esporta Giacenze")
        self.btn_export_giacenze.setProperty('class', 'success')
        self.btn_export_giacenze.clicked.connect(self._esporta_giacenze)

        # Toggle reale ↔ fittizio. Classe `secondary` (blu): nel QSS globale
        # esistono solo success/warning/danger/secondary; "primary" non c'è e
        # ricadeva nel default Qt fuori standard.
        self.btn_toggle_fittizio = QPushButton("📋 Mostra Fittizio")
        self.btn_toggle_fittizio.setProperty('class', 'secondary')
        self.btn_toggle_fittizio.clicked.connect(self._toggle_fittizio)

        top_layout = self.layout().itemAt(0).layout()
        top_layout.insertWidget(4, self.btn_export_tutti_mov)
        top_layout.insertWidget(5, self.btn_export_giacenze)
        top_layout.insertWidget(6, self.btn_toggle_fittizio)

        self.vista.doubleClicked.connect(self._on_doppio_click)
        self.aggiorna_dati()

    def _tabella(self) -> str:
        """Nome SQL della tabella di lettura corrente (reale o fittizio).
        Le scritture manuali vanno sempre nel reale e sono bloccate in
        modalità fittizio (vedi DialogRegistroProdotto)."""
        return "registro_magazzino_fittizio" if self.mostra_fittizio \
            else "registro_magazzino"

    def _toggle_fittizio(self):
        self.mostra_fittizio = not self.mostra_fittizio
        if self.mostra_fittizio:
            self.btn_toggle_fittizio.setText("📦 Mostra Reale")
        else:
            self.btn_toggle_fittizio.setText("📋 Mostra Fittizio")
        self.aggiorna_dati()

    def _resolve_filter_ids(self) -> list[int]:
        """Risolve azienda_filter (nome) in lista di azienda_id da considerare.

        Include anche gli alias che mappano sull'azienda target (es. quando il
        filtro è 'Messina Alfio', include anche 'Deflorio Ciccopinto').
        Ritorna [] se filtro vuoto o azienda inesistente nel DB.
        """
        if not self.azienda_filter:
            return []
        from config import WAREHOUSE_ALIASES
        names = [self.azienda_filter]
        for alias, canonical in WAREHOUSE_ALIASES.items():
            if canonical.lower() == self.azienda_filter.lower():
                names.append(alias)
        ids = []
        with self.engine.connect() as conn:
            for n in names:
                row = conn.execute(
                    text("SELECT id FROM aziende WHERE LOWER(TRIM(nome)) = LOWER(TRIM(:n))"),
                    {"n": n}
                ).first()
                if row:
                    ids.append(row[0])
        return ids

    def _where_azienda_clause(self, table_alias: str) -> str:
        """Costruisce il pezzo SQL 'AND <alias>.azienda_id IN (...)'.
        Stringa vuota se nessun filtro. Usa int() coercion per safety."""
        if not self.azienda_ids:
            return ""
        ids_str = ",".join(str(int(i)) for i in self.azienda_ids)
        return f"AND {table_alias}.azienda_id IN ({ids_str})"

    def aggiorna_dati(self):
        suffisso = " — Fittizio" if self.mostra_fittizio else ""
        giacenza_label = f"Giacenza ({self.azienda_filter}){suffisso}"
        clausola_az = self._where_azienda_clause("rm")  # es. "AND rm.azienda_id IN (1,5)"
        tabella = self._tabella()
        query = f"""
            SELECT p.id, p.nome_prodotto AS "Nome Prodotto",
                   ROUND(COALESCE((
                       SELECT SUM(CASE WHEN rm.tipo_movimento = 'CARICO' THEN rm.quantita ELSE -rm.quantita END)
                       FROM {tabella} rm
                       WHERE rm.prodotto_id = p.id {clausola_az}
                   ), 0), 4) AS "{giacenza_label}",
                   p.categoria AS "Categoria", p.unita_misura AS "Unità", p.numero_registrazione AS "N. Registrazione",
                   p.sostanza_attiva AS "Sostanza Attiva", p.bio_convenzionale AS "Bio/Conv", p.blacklist AS "Blacklist",
                   p.avversita AS "Avversità", p.titolo_n, p.titolo_p, p.titolo_k, p.phi_giorni, p.trattamenti_max, p.intervallo_min_tratt, p.min_sostanza, p.max_sostanza, p.qta_acqua
            FROM prodotti p ORDER BY p.nome_prodotto
        """
        self.esegui_query(query, self.engine)
        self._nascondi_colonne()
        self.vista.resizeColumnsToContents()

    def _esporta_tutti_movimenti(self):
        # Nome file di default include il magazzino, se filtrato.
        # Sanitizziamo i caratteri illegali su Windows (\/:*?"<>|) sostituendoli
        # con underscore. Senza, un'azienda nominata "Messina: Alfio" rendeva
        # il dialogo di salvataggio fallimentare oppure salvava un file con
        # nome corrotto in modo confuso per l'utente.
        import re
        if self.azienda_filter:
            safe_name = re.sub(r'[\\/:*?"<>|]', '_', self.azienda_filter).replace(' ', '_')
            default_name = f"Movimenti_{safe_name}.xlsx"
        else:
            default_name = "Movimenti_Magazzino.xlsx"
        f_p, _ = QFileDialog.getSaveFileName(self, "Esporta Movimenti", default_name, "Excel (*.xlsx)")
        if not f_p:
            return
        if not f_p.lower().endswith('.xlsx'):
            f_p += '.xlsx'

        try:
            # Filtra per i magazzini di questa veduta
            clausola_az_where = ""
            if self.azienda_ids:
                ids_str = ",".join(str(int(i)) for i in self.azienda_ids)
                clausola_az_where = f"WHERE rm.azienda_id IN ({ids_str})"
            with self.engine.connect() as conn:
                query = text(f"""
                    SELECT
                        rm.id AS _id,
                        COALESCE(az_orig.nome, az_wh.nome, 'Generale') AS "Azienda",
                        rm.azienda_id AS _azienda_id,
                        rm.prodotto_id AS _prodotto_id,
                        p.nome_prodotto AS "Prodotto",
                        p.numero_registrazione AS "N. Registrazione",
                        p.sostanza_attiva AS "Sostanza Attiva",
                        rm.n_ddt AS "N. DDT",
                        rm.fornitore AS "Fornitore",
                        rm.data_movimento AS "Data",
                        rm.tipo_movimento AS "Tipo",
                        rm.quantita AS _qta_raw,
                        p.unita_misura AS _um
                    FROM {self._tabella()} rm
                    JOIN prodotti p ON p.id = rm.prodotto_id
                    LEFT JOIN aziende az_wh ON az_wh.id = rm.azienda_id
                    LEFT JOIN aziende az_orig ON az_orig.id = rm.azienda_id_origine
                    {clausola_az_where}
                    ORDER BY rm.data_movimento ASC, rm.id ASC
                """)
                df = pd.read_sql(query, conn)

            if df.empty:
                QMessageBox.information(self, "Nessun dato", "Nessun movimento da esportare.")
                return

            def normalizza_tipo(t):
                t_upper = str(t).strip().upper() if t else ""
                if "SCARICO" in t_upper or "USCITA" in t_upper:
                    return "Scarico"
                return "Carico"

            df["Carico/Scarico"] = df["Tipo"].apply(normalizza_tipo)
            df["Quantità"] = df.apply(
                lambda r: -abs(r["_qta_raw"]) if r["Carico/Scarico"] == "Scarico" else abs(r["_qta_raw"]),
                axis=1
            )

            df["Giacenza"] = (
                df.groupby(["_azienda_id", "_prodotto_id"])["Quantità"]
                  .cumsum()
                  .round(4)
            )

            df["Unità"] = df["_um"]

            colonne_finali = [
                "Prodotto", "N. Registrazione", "Sostanza Attiva",
                "N. DDT", "Fornitore", "Data", "Carico/Scarico",
                "Quantità", "Unità", "Giacenza", "Azienda"
            ]
            df = df[colonne_finali]

            df = df.sort_values(
                by=["Azienda", "Data", "Prodotto"],
                ascending=[True, False, True]
            ).reset_index(drop=True)

            with pd.ExcelWriter(f_p, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name="Movimenti Magazzino", index=False)

                ws = writer.sheets["Movimenti Magazzino"]
                from openpyxl.styles import Font, PatternFill, Alignment

                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
                header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_align

                fill_carico = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
                fill_scarico = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

                col_tipo_idx = colonne_finali.index("Carico/Scarico") + 1
                for row_idx in range(2, ws.max_row + 1):
                    tipo_val = ws.cell(row=row_idx, column=col_tipo_idx).value
                    fill_riga = fill_scarico if tipo_val == "Scarico" else fill_carico
                    for col_idx in range(1, len(colonne_finali) + 1):
                        ws.cell(row=row_idx, column=col_idx).fill = fill_riga

                for col_idx, column in enumerate(ws.columns, start=1):
                    max_length = max(
                        (len(str(cell.value)) for cell in column if cell.value is not None),
                        default=10
                    )
                    col_letter = ws.cell(row=1, column=col_idx).column_letter
                    ws.column_dimensions[col_letter].width = min(max_length + 2, 35)

                ws.row_dimensions[1].height = 30
                ws.freeze_panes = "A2"

            QMessageBox.information(self, "Esportazione", f"File Excel salvato correttamente:\n{f_p}")
        except ImportError:
            QMessageBox.critical(self, "Errore", "Per esportare in Excel servono le librerie pandas e openpyxl.\nInstallale con: pip install pandas openpyxl")
        except PermissionError:
            # Caso comune: l'utente ha già aperto il file in Excel/LibreOffice.
            # Senza handle dedicato, l'errore generico era misterioso.
            QMessageBox.critical(
                self, "File in uso",
                f"Impossibile scrivere il file:\n{f_p}\n\n"
                "È aperto in Excel o LibreOffice. Chiudilo e riprova.",
            )
        except Exception as e:
            log.exception("Esportazione movimenti Excel fallita")
            QMessageBox.critical(self, "Errore", f"Esportazione fallita:\n{str(e)}")

    def _esporta_giacenze(self):
        """Esporta in Excel la giacenza prodotto-per-prodotto.

        Rispetta:
        - Filtro azienda (`azienda_ids`): include solo i magazzini selezionati
          + i loro alias (vedi `_resolve_filter_ids`).
        - Toggle reale/fittizio (`mostra_fittizio`): legge dalla tabella
          corrente, così l'utente esporta lo snapshot che sta guardando.
        - Prodotti in blacklist sono inclusi con flag dedicato per coerenza
          con la vista (non esclusi: l'utente potrebbe averne in giacenza
          da prima della messa in blacklist).
        """
        import re
        from datetime import datetime

        suffisso = "_fittizio" if self.mostra_fittizio else ""
        if self.azienda_filter:
            safe_name = re.sub(r'[\\/:*?"<>|]', '_', self.azienda_filter).replace(' ', '_')
            default_name = f"Giacenze_{safe_name}{suffisso}_{datetime.now():%Y-%m-%d}.xlsx"
        else:
            default_name = f"Giacenze_Magazzino{suffisso}_{datetime.now():%Y-%m-%d}.xlsx"

        f_p, _ = QFileDialog.getSaveFileName(self, "Esporta Giacenze", default_name, "Excel (*.xlsx)")
        if not f_p:
            return
        if not f_p.lower().endswith('.xlsx'):
            f_p += '.xlsx'

        tabella = self._tabella()
        clausola_az = self._where_azienda_clause("rm")

        try:
            import pandas as pd

            with self.engine.connect() as conn:
                # Stesso calcolo di `aggiorna_dati`: SUM(CARICO) - SUM(SCARICO)
                # per prodotto. Aggiungiamo anche conteggi N. movimenti +
                # ultimo movimento (data) per dare contesto all'inventario.
                # Lasciato fuori `qta_acqua` / `titolo_n,p,k` che servono solo
                # alla UI per i dettagli prodotto.
                query = text(f"""
                    SELECT
                        p.nome_prodotto AS "Prodotto",
                        p.unita_misura AS "Unità Misura",
                        ROUND(COALESCE((
                            SELECT SUM(CASE WHEN rm.tipo_movimento = 'CARICO'
                                       THEN rm.quantita ELSE -rm.quantita END)
                            FROM {tabella} rm
                            WHERE rm.prodotto_id = p.id {clausola_az}
                        ), 0), 4) AS "Giacenza",
                        (SELECT COUNT(*) FROM {tabella} rm
                          WHERE rm.prodotto_id = p.id {clausola_az}) AS "N. Movimenti",
                        (SELECT MAX(rm.data_movimento) FROM {tabella} rm
                          WHERE rm.prodotto_id = p.id {clausola_az}) AS "Ultimo Movimento",
                        p.categoria AS "Categoria",
                        p.numero_registrazione AS "N. Registrazione",
                        p.sostanza_attiva AS "Sostanza Attiva",
                        p.bio_convenzionale AS "Bio/Conv",
                        CASE WHEN LOWER(TRIM(COALESCE(p.blacklist,''))) = 'si'
                             THEN 'BLACKLIST' ELSE '' END AS "Blacklist",
                        p.avversita AS "Avversità",
                        p.phi_giorni AS "PHI (giorni)",
                        p.trattamenti_max AS "Max Trattamenti",
                        p.intervallo_min_tratt AS "Intervallo Min (gg)",
                        p.min_sostanza AS "Min Etichetta",
                        p.max_sostanza AS "Max Etichetta"
                    FROM prodotti p
                    ORDER BY p.nome_prodotto
                """)
                df = pd.read_sql(query, conn)

            if df.empty:
                QMessageBox.information(self, "Nessun dato",
                                        "Nessun prodotto in anagrafica.")
                return

            with pd.ExcelWriter(f_p, engine='openpyxl') as writer:
                # Sheet name include il magazzino se filtrato — utile quando
                # l'utente esporta più magazzini e poi li accorpa in un solo workbook.
                sheet_name = (self.azienda_filter or "Magazzino")[:31]  # 31 = limite Excel
                df.to_excel(writer, sheet_name=sheet_name, index=False)

                ws = writer.sheets[sheet_name]
                from openpyxl.styles import Font, PatternFill, Alignment

                # Header verde brand
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
                header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_align

                # Evidenzia righe blacklist (rosso chiaro), giacenza negativa
                # (rosso più scuro) e righe a zero (grigio chiaro).
                fill_blacklist = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
                fill_negativa = PatternFill(start_color="EF9A9A", end_color="EF9A9A", fill_type="solid")
                fill_zero = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
                col_giacenza = df.columns.get_loc("Giacenza") + 1
                col_blacklist = df.columns.get_loc("Blacklist") + 1
                for row_idx in range(2, ws.max_row + 1):
                    bl = ws.cell(row=row_idx, column=col_blacklist).value
                    gi = ws.cell(row=row_idx, column=col_giacenza).value
                    fill = None
                    if bl == "BLACKLIST":
                        fill = fill_blacklist
                    elif isinstance(gi, (int, float)) and gi < 0:
                        fill = fill_negativa
                    elif isinstance(gi, (int, float)) and gi == 0:
                        fill = fill_zero
                    if fill:
                        for col_idx in range(1, len(df.columns) + 1):
                            ws.cell(row=row_idx, column=col_idx).fill = fill

                # Auto-fit colonne (con clamp).
                for col_idx, column in enumerate(ws.columns, start=1):
                    max_length = max(
                        (len(str(cell.value)) for cell in column if cell.value is not None),
                        default=10,
                    )
                    col_letter = ws.cell(row=1, column=col_idx).column_letter
                    ws.column_dimensions[col_letter].width = min(max_length + 2, 35)
                ws.row_dimensions[1].height = 30
                ws.freeze_panes = "A2"

            QMessageBox.information(self, "Esportazione",
                                    f"File Excel salvato correttamente:\n{f_p}")
        except ImportError:
            QMessageBox.critical(self, "Errore",
                                 "Per esportare in Excel servono le librerie pandas e openpyxl.\n"
                                 "Installale con: pip install pandas openpyxl")
        except PermissionError:
            QMessageBox.critical(
                self, "File in uso",
                f"Impossibile scrivere il file:\n{f_p}\n\n"
                "È aperto in Excel o LibreOffice. Chiudilo e riprova.",
            )
        except Exception as e:
            log.exception("Esportazione giacenze Excel fallita")
            QMessageBox.critical(self, "Errore", f"Esportazione fallita:\n{str(e)}")

    def _on_doppio_click(self, index):
        # --- FIX: Recupero dati compatibile con QStandardItemModel ---
        item_col0 = self.modello.item(index.row(), 0)
        dati = item_col0.data(Qt.ItemDataRole.UserRole)

        if not dati:
            return

        # Creiamo un "finto record" per non dover riscrivere il resto della funzione
        class DummyRecord:
            def __init__(self, d): self.d = d
            def value(self, k): return self.d.get(k)

        r = DummyRecord(dati)
        # -------------------------------------------------------------

        DialogRegistroProdotto(
            self.engine, self.db, r.value("id"),
            r.value("Nome Prodotto"), r.value("Unità"), self,
            azienda_filter=self.azienda_filter,
            azienda_ids=self.azienda_ids,
            mostra_fittizio=self.mostra_fittizio,
        ).exec()
        self.aggiorna_dati()

    def apri_dialog_nuovo(self):
        if DialogProdotto(self.engine, parent=self).exec(): self.aggiorna_dati()

    def apri_dialog_modifica(self, riga):
        # `unita_carico` non viene esposto come colonna della tabella prodotti
        # (vedi aggiorna_dati): lo leggiamo direttamente dal DB così il combo
        # del dialog si apre con la selezione corretta.
        prod_id = riga.value("id")
        with self.engine.connect() as conn:
            unita_carico = conn.execute(text(
                "SELECT unita_carico FROM prodotti WHERE id = :pid"
            ), {"pid": prod_id}).scalar()
        dati = {
            "id": prod_id, "nome_prodotto": riga.value("Nome Prodotto"), "categoria": riga.value("Categoria"),
            "numero_registrazione": riga.value("N. Registrazione"), "sostanza_attiva": riga.value("Sostanza Attiva"),
            "bio_convenzionale": riga.value("Bio/Conv"), "blacklist": riga.value("Blacklist"), "avversita": riga.value("Avversità"),
            "unita_misura": riga.value("Unità"), "unita_carico": unita_carico,
            "titolo_n": riga.value("titolo_n"), "titolo_p": riga.value("titolo_p"), "titolo_k": riga.value("titolo_k"),
            "phi_giorni": riga.value("phi_giorni"), "trattamenti_max": riga.value("trattamenti_max"), "intervallo_min_tratt": riga.value("intervallo_min_tratt"),
            "min_sostanza": riga.value("min_sostanza"), "max_sostanza": riga.value("max_sostanza"), "qta_acqua": riga.value("qta_acqua")
        }
        if DialogProdotto(self.engine, dati=dati, parent=self).exec(): self.aggiorna_dati()

    def elimina_record(self, riga):
        if QMessageBox.question(self, "Elimina", f"Eliminare {riga.value('Nome Prodotto')}?") == QMessageBox.StandardButton.Yes:
            with self.engine.begin() as conn: conn.execute(text("DELETE FROM prodotti WHERE id=:id"), {"id": riga.value("id")})
            self.aggiorna_dati()
